# -*- coding: utf-8 -*-
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import os

def create_element(name):
    return OxmlElement(name)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_shading(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color_hex)
    tcPr.append(shd)

def set_cell_border(cell, top="D3D3D3", bottom="D3D3D3", left="none", right="none", sz="4"):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    borders = {'top': top, 'bottom': bottom, 'left': left, 'right': right}
    for border_name, border_color in borders.items():
        border = OxmlElement(f'w:{border_name}')
        if border_color == 'none':
            border.set(qn('w:val'), 'none')
        else:
            border.set(qn('w:val'), 'single')
            border.set(qn('w:sz'), sz)
            border.set(qn('w:space'), '0')
            border.set(qn('w:color'), border_color)
        tcBorders.append(border)
    tcPr.append(tcBorders)

def add_code_box(doc, code_text, caption=None):
    if caption:
        p_cap = doc.add_paragraph()
        p_cap.paragraph_format.space_before = Pt(6)
        p_cap.paragraph_format.space_after = Pt(2)
        p_cap.paragraph_format.keep_with_next = True
        run_cap = p_cap.add_run(caption)
        run_cap.font.name = "Times New Roman"
        run_cap.font.size = Pt(11)
        run_cap.font.bold = True
        run_cap.font.italic = True
        run_cap.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.rows[0].cells[0]
    set_cell_shading(cell, "F8F9FA")
    set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
    
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for b_name in ['top', 'left', 'bottom', 'right']:
        b = OxmlElement(f'w:{b_name}')
        b.set(qn('w:val'), 'single')
        b.set(qn('w:sz'), '6')
        b.set(qn('w:space'), '0')
        b.set(qn('w:color'), '000000')
        tcBorders.append(b)
    tcPr.append(tcBorders)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(code_text.strip())
    run.font.name = "Consolas"
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor(0x20, 0x20, 0x20)

    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(6)

def add_callout(doc, text, title=None, alert_type="NOTE"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.rows[0].cells[0]
    
    colors = {
        "NOTE": ("F0F4F8", "1F497D"),
        "WARNING": ("FFF8E7", "B78103"),
        "IMPORTANT": ("EBF5FB", "2980B9"),
        "FIX": ("E8F8F5", "117864")
    }
    bg_color, border_color = colors.get(alert_type, ("F8F9FA", "555555"))
    
    set_cell_shading(cell, bg_color)
    set_cell_margins(cell, top=100, bottom=100, left=180, right=140)
    
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    left_b = OxmlElement('w:left')
    left_b.set(qn('w:val'), 'single')
    left_b.set(qn('w:sz'), '24')
    left_b.set(qn('w:space'), '0')
    left_b.set(qn('w:color'), border_color)
    tcBorders.append(left_b)
    for b_name in ['top', 'bottom', 'right']:
        b = OxmlElement(f'w:{b_name}')
        b.set(qn('w:val'), 'none')
        tcBorders.append(b)
    tcPr.append(tcBorders)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.2
    
    if title:
        run_title = p.add_run(f"[{title}]\n")
        run_title.font.name = "Times New Roman"
        run_title.font.size = Pt(11.5)
        run_title.font.bold = True
        run_title.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)
        
    run_text = p.add_run(text)
    run_text.font.name = "Times New Roman"
    run_text.font.size = Pt(11)
    run_text.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(6)

def add_styled_table(doc, headers, data, col_widths=None):
    tbl = doc.add_table(rows=len(data) + 1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    hdr_row = tbl.rows[0]
    for idx, h_text in enumerate(headers):
        cell = hdr_row.cells[idx]
        set_cell_shading(cell, "F2F2F2")
        set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
        set_cell_border(cell, top="000000", bottom="000000", left="none", right="none", sz="8")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(h_text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(10.5)
        run.font.bold = True
        
    for r_idx, row_data in enumerate(data):
        row = tbl.rows[r_idx + 1]
        bg = "FFFFFF" if r_idx % 2 == 0 else "FBFBFB"
        for c_idx, val in enumerate(row_data):
            cell = row.cells[c_idx]
            set_cell_shading(cell, bg)
            set_cell_margins(cell, top=70, bottom=70, left=100, right=100)
            is_last = (r_idx == len(data) - 1)
            b_bottom = "000000" if is_last else "E0E0E0"
            b_sz = "8" if is_last else "4"
            set_cell_border(cell, top="none", bottom=b_bottom, left="none", right="none", sz=b_sz)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            run = p.add_run(str(val))
            run.font.name = "Times New Roman"
            run.font.size = Pt(10)

    if col_widths:
        for r in tbl.rows:
            for c_idx, w in enumerate(col_widths):
                r.cells[c_idx].width = Inches(w)
                
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(6)

def set_para_font(p, text="", bold=False, italic=False, size=13, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=0, space_after=6, line_spacing=1.25):
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    if text:
        r = p.add_run(text)
        r.font.name = "Times New Roman"
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
        return r
    return None

def add_warning_box(doc, text, title="NỘI DUNG CŨ CẦN XÓA / SỬA (FIND & REMOVE)"):
    add_callout(doc, text, title=title, alert_type="WARNING")

def add_fix_box(doc, text, title="NỘI DUNG MỚI CẦN COPY & DÁN ĐÈ VÀO"):
    add_callout(doc, text, title=title, alert_type="FIX")

def add_location_box(doc, chapter, section_num, section_title, target_desc, search_keyword):
    headers = ["Thuộc tính", "Thông tin đối chiếu trong Báo cáo gốc"]
    data = [
        ["Chương / Phần", chapter],
        ["Đề mục cụ thể", f"Mục {section_num}: {section_title}"],
        ["Vị trí chính xác", target_desc],
        ["Từ khóa tìm kiếm (Ctrl+F)", search_keyword]
    ]
    add_styled_table(doc, headers, data, col_widths=[2.0, 5.2])

print("Helper definitions loaded successfully.")
