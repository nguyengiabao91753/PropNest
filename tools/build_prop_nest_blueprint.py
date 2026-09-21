from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from pathlib import Path


OUT = Path(__file__).resolve().parents[1] / "docs" / "PropNest_Architecture_and_Implementation_Blueprint.docx"

NAVY = "17365D"
PALE_BLUE = "EAF2F8"
LIGHT_GRAY = "F3F5F7"
BORDER = "D9D9D9"
BLACK = RGBColor(0, 0, 0)


def set_font(run, size=11, bold=None, italic=None, color=BLACK, name="Arial"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=BORDER):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{side}")
        edge = borders.find(tag)
        if edge is None:
            edge = OxmlElement(f"w:{side}")
            borders.append(edge)
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "6")
        edge.set(qn("w:space"), "0")
        edge.set(qn("w:color"), color)


def set_cell_margins(cell, top=100, start=110, bottom=100, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:cantSplit")
    tr_pr.append(node)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Trang ")
    set_font(run, 9, color=BLACK)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def add_text(doc, text="", bold_lead=None, space_after=5, keep=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.12
    if keep:
        p.paragraph_format.keep_with_next = True
    if bold_lead and text.startswith(bold_lead):
        lead = p.add_run(bold_lead)
        set_font(lead, 11, bold=True)
        body = p.add_run(text[len(bold_lead):])
        set_font(body, 11)
    else:
        run = p.add_run(text)
        set_font(run, 11)
    return p


def add_code(doc, lines):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.left_indent = Inches(0.18)
    for i, line in enumerate(lines):
        if i:
            p.add_run("\n")
        run = p.add_run(line)
        set_font(run, 9.5, name="Consolas")
    return p


def add_bullets(doc, items, level=0):
    for item in items:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.08
        run = p.add_run(item)
        set_font(run, 11)


def heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.space_before = Pt(13 if level == 1 else 9)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.keep_with_next = True
    for r in p.runs:
        set_font(r, 16 if level == 1 else 12.5, bold=True)
    return p


def add_table(doc, headers, rows, widths=None, font_size=9.4):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    hdr = table.rows[0]
    repeat_table_header(hdr)
    prevent_row_split(hdr)
    for i, value in enumerate(headers):
        cell = hdr.cells[i]
        cell.text = ""
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, NAVY)
        set_cell_border(cell)
        set_cell_margins(cell)
        run = cell.paragraphs[0].add_run(value)
        set_font(run, font_size, bold=True, color=RGBColor(255, 255, 255))
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        if widths:
            cell.width = Inches(widths[i])
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        prevent_row_split(table.rows[-1])
        for i, value in enumerate(values):
            cell = cells[i]
            cell.text = ""
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_shading(cell, "FFFFFF" if row_index % 2 == 0 else PALE_BLUE)
            set_cell_border(cell)
            set_cell_margins(cell)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            run = p.add_run(str(value))
            set_font(run, font_size)
            if widths:
                cell.width = Inches(widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_table_caption(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(text)
    set_font(run, 10, bold=True)
    return p


def configure_document(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.70)
    section.right_margin = Inches(0.70)
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(11)
    normal.font.color.rgb = BLACK

    for name in ("Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.color.rgb = BLACK

    doc.styles["Title"].font.size = Pt(24)
    doc.styles["Heading 1"].font.size = Pt(16)
    doc.styles["Heading 2"].font.size = Pt(12.5)
    doc.styles["Heading 3"].font.size = Pt(11)

    footer = section.footer
    add_page_number(footer.paragraphs[0])


def title_page(doc):
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(110)
    p.paragraph_format.space_after = Pt(12)
    run = p.add_run("Định hướng kiến trúc và triển khai hệ thống PropNest")
    set_font(run, 24, bold=True)

    p = doc.add_paragraph(style="Subtitle")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(26)
    run = p.add_run("Bản thiết kế kỹ thuật cho dự án hệ thống đăng tin bất động sản")
    set_font(run, 13)

    add_text(doc, "Tài liệu này chốt hướng phát triển cho nhóm theo mô hình Modular Monolith, sử dụng ASP.NET Core, SQL Server và Saga Orchestrator. Mục tiêu là dùng các feature nghiệp vụ để chứng minh các kỹ thuật .NET quan trọng, đồng thời giữ kiến trúc đủ mở để tích hợp AME tự động kiểm duyệt sau khi phần lõi ổn định.", space_after=12)
    add_text(doc, "Phạm vi tài liệu gồm kiến trúc, feature backlog, thư viện, kỹ thuật .NET, database, workflow, API chủ đạo, kiểm thử và lộ trình tích hợp AME.", space_after=16)

    add_table_caption(doc, "Quyết định kiến trúc")
    add_table(doc,
              ["Nội dung", "Quyết định"],
              [
                  ["Kiểu hệ thống", "Modular Monolith gồm API và Workflow Worker; cùng một codebase, SQL Server và RabbitMQ."],
                  ["Điều phối giao dịch", "Saga Orchestration qua MassTransit state machine; không dùng choreography cho workflow thanh toán chính."],
                  ["Tính nhất quán", "Idempotency request, optimistic concurrency, transactional outbox, inbox deduplication và compensating action."],
                  ["Kiểm duyệt", "Manual moderation ở bản lõi; AME FastAPI và Ollama là phase sau, kết nối bất đồng bộ bằng event."],
              ], widths=[1.65, 5.35], font_size=10)
    doc.add_page_break()


def add_overview(doc):
    heading(doc, "1 Mục tiêu và phạm vi", 1)
    add_text(doc, "PropNest là hệ thống cho phép người bán tạo và quản lý tin bất động sản, mua gói hiển thị, theo dõi lịch sử thay đổi và đưa tin vào quy trình kiểm duyệt. Trọng tâm học phần là chất lượng thiết kế và khả năng giải thích cơ chế ASP.NET Core, không phải mô phỏng toàn bộ quy mô enterprise.")
    heading(doc, "1 1 Mục tiêu học phần được thể hiện bằng nghiệp vụ", 2)
    add_table(doc,
              ["Mục tiêu", "Nghiệp vụ minh họa", "Bằng chứng demo"],
              [
                  ["ASP.NET Core pipeline", "Tạo tin và mua gói", "Thứ tự middleware, correlation ID, xử lý exception nhất quán."],
                  ["Dependency injection", "Listing, Wallet, Delta engine", "Giải thích scoped DbContext, scoped service và singleton configuration."],
                  ["EF Core và migration", "Schema listing, ví, history, outbox", "Migration tạo database, quan hệ, index và rowversion."],
                  ["Authentication và authorization", "Seller sửa tin, Moderator duyệt tin", "JWT, role policy và resource based authorization."],
                  ["Distributed workflow", "Mua và nâng cấp gói VIP", "Saga state, retry, compensation refund, idempotency."],
                  ["Observability", "Tất cả request và workflow", "Structured log, health check, correlation ID và trạng thái workflow."],
              ], widths=[1.55, 2.20, 3.25])
    heading(doc, "1 2 Phạm vi phát hành", 2)
    add_bullets(doc, [
        "MVP bắt buộc: xác thực, CRUD listing, vòng đời tin, duyệt thủ công, history delta, ví và mua gói với workflow theo dõi được.",
        "Nâng cao: RabbitMQ, MassTransit Saga, transactional outbox, Quartz expiration job, integration tests và dashboard workflow tối giản.",
        "AME chỉ bắt đầu sau khi workflow core có test và demo ổn định. Ảnh và video moderation là phase tiếp theo sau AME text and price.",
    ])
    heading(doc, "1 3 Các trạng thái quan trọng", 2)
    add_table(doc,
              ["Đối tượng", "Trạng thái đề xuất", "Quy tắc chính"],
              [
                  ["Listing", "Draft, PendingPayment, PendingModeration, Published, Rejected, Hidden, Expired, Deleted", "Chỉ Published được public. Sửa nội dung Published có thể quay lại PendingModeration."],
                  ["Workflow", "Started, DeductingWallet, WalletDeducted, UpgradingListing, RecordingHistory, Completed, Compensating, Compensated, Failed", "Workflow là nguồn trạng thái cho giao dịch mua gói."],
                  ["Moderation", "NotRequested, Pending, Approved, ManualReview, Rejected", "Decision tách khỏi ListingStatus để lưu được lý do và nguồn quyết định."],
                  ["Outbox", "Pending, Processing, Published, Failed", "Message được phát sau khi transaction database đã commit."],
              ], widths=[1.35, 2.9, 2.75])


def add_architecture(doc):
    heading(doc, "2 Kiến trúc hệ thống", 1)
    add_text(doc, "Kiến trúc được chọn là Modular Monolith theo Clean Architecture. Các module dùng chung database nhưng chỉ giao tiếp qua application contracts và domain events. API không sở hữu state Saga; state và consumers nằm trong Workflow Worker để workflow vẫn tiếp tục được khi API restart.")
    heading(doc, "2 1 Sơ đồ triển khai", 2)
    add_code(doc, [
        "Frontend and Swagger",
        "        | HTTPS and JWT",
        "PropNest Api  ->  Application  ->  Domain",
        "        |                 |",
        "        |                 -> Infrastructure -> SQL Server",
        "        |                                  -> Outbox",
        "        |                                  -> RabbitMQ",
        "        |",
        "Workflow Worker -> MassTransit Saga and Consumers -> RabbitMQ and SQL Server",
        "        |",
        "Phase later: AME FastAPI -> Rule Engine and Ollama -> AmeModerationResult event",
    ])
    add_table_caption(doc, "Trách nhiệm của từng tiến trình")
    add_table(doc,
              ["Thành phần", "Trách nhiệm", "Không làm"],
              [
                  ["PropNest Api", "HTTP endpoint, JWT, validation, idempotency entry point, query status workflow.", "Không chờ toàn bộ Saga; không giữ Saga state trong memory; không chứa business logic dài."],
                  ["Application", "Use cases, commands and queries, interfaces, transaction boundary, publish domain integration event.", "Không biết SQL Server, RabbitMQ hoặc HTTP framework cụ thể."],
                  ["Domain", "Entity, enum, invariant, trạng thái listing và ví, domain exception.", "Không phụ thuộc EF Core, controller, queue hoặc LLM."],
                  ["Infrastructure", "EF Core, repositories, outbox dispatcher, JWT implementation, Serilog, SQL persistence.", "Không quyết định quy tắc business từ controller."],
                  ["Workflow Worker", "Saga state machine, consumers, retry, compensation, inbox deduplication.", "Không expose HTTP trực tiếp cho Frontend."],
                  ["AME service", "Rule checks, local LLM analysis, decision aggregation và trả moderation result.", "Không tự publish listing trực tiếp vào database PropNest."],
              ], widths=[1.25, 3.25, 2.50])
    heading(doc, "2 2 Dependency rule", 2)
    add_code(doc, [
        "PropNest Api -> Application, Infrastructure",
        "Workflow Worker -> Application, Infrastructure",
        "Infrastructure -> Application, Domain",
        "Application -> Domain",
        "Domain -> no project dependency",
    ])
    heading(doc, "2 3 Cấu trúc solution", 2)
    add_code(doc, [
        "src/PropNest.Api                 controllers, middleware, filters, OpenAPI",
        "src/PropNest.Application         feature commands, queries, DTOs, validators, interfaces",
        "src/PropNest.Domain              entities, value objects, domain events, enums",
        "src/PropNest.Infrastructure      EF Core, repositories, identity, outbox, messaging",
        "src/PropNest.Workflow.Worker     MassTransit consumers, saga state machine, Quartz jobs",
        "tests/PropNest.Domain.Tests      pure business and Delta engine tests",
        "tests/PropNest.Application.Tests command handler and policy tests",
        "tests/PropNest.IntegrationTests  API, SQL Server and RabbitMQ integration tests",
        "services/ame                     FastAPI and Ollama prototype after core completion",
    ])
    heading(doc, "2 4 Middleware pipeline", 2)
    add_code(doc, [
        "ExceptionHandlingMiddleware",
        "-> CorrelationIdMiddleware",
        "-> Serilog request logging",
        "-> HTTPS and CORS",
        "-> Authentication",
        "-> Authorization",
        "-> Endpoint or Controller",
    ])
    add_text(doc, "Khi demo, đảo CorrelationIdMiddleware xuống sau request logging để cho thấy log đầu request không còn correlation ID. Đây là minh chứng trực tiếp rằng thứ tự middleware làm thay đổi kết quả.")


def add_dotnet_practices(doc):
    heading(doc, "3 Kỹ thuật NET và thư viện theo feature", 1)
    add_text(doc, "Không cần thêm tất cả thư viện ngay từ đầu. Cài package khi feature tương ứng bắt đầu, lock version theo Target Framework của solution và ghi lại quyết định trong README.")
    add_table(doc,
              ["Kỹ thuật NET", "Feature áp dụng", "Cách dùng cụ thể"],
              [
                  ["Minimal hosting và endpoint routing", "Toàn hệ thống", "Program.cs đăng ký service, middleware, authentication, authorization, OpenAPI và health checks."],
                  ["Model binding and FluentValidation", "Create and Update Listing, top up, purchase package", "Request DTO nhận từ HTTP; validator trả lỗi 400 ProblemDetails trước khi handler chạy."],
                  ["Dependency injection", "Listing, Wallet, History, Saga", "DbContext, repositories và application services scoped; options and stateless infrastructure services singleton khi thread safe."],
                  ["Exception middleware", "Toàn hệ thống", "Chuyển DomainException, NotFoundException, ConcurrencyException thành ProblemDetails chuẩn."],
                  ["Action filter hoặc endpoint filter", "API write endpoints", "Gắn audit metadata hoặc kiểm tra Idempotency-Key ở một chỗ. Không dùng filter để che business workflow."],
                  ["JWT and policy authorization", "Login, Listing, Moderation, Wallet", "Seller sở hữu resource; Moderator hoặc Admin duyệt; policy tách khỏi controller logic."],
                  ["EF Core migration and Fluent API", "Database core", "Map enum, decimal precision, foreign key, unique index, RowVersion và query filter nếu dùng soft delete."],
                  ["Optimistic concurrency", "Update Listing and Wallet", "Listings và Wallets có RowVersion; client gửi ETag or version để phát hiện ghi đè."],
                  ["BackgroundService and Quartz", "Outbox, expiry, AME retry", "Worker tạo scope mới cho mỗi job, không inject DbContext vào singleton."],
                  ["Options pattern", "JWT, RabbitMQ, package policy, AME endpoint", "Bind strongly typed settings; validate configuration khi startup."],
                  ["Structured logging and health checks", "API and Worker", "Serilog chứa correlation ID and workflow ID; /health kiểm tra SQL Server and RabbitMQ."],
                  ["CancellationToken and async", "API and consumers", "Truyền token đến EF, HTTP AME and message consumer; không block thread bằng .Result."],
              ], widths=[1.65, 2.15, 3.20], font_size=9.1)
    heading(doc, "3 1 Thư viện đề xuất", 2)
    add_table(doc,
              ["Package hoặc công cụ", "Dùng cho", "Feature liên quan", "Ghi chú"],
              [
                  ["Microsoft.EntityFrameworkCore.SqlServer and Tools", "ORM, migration", "Tất cả dữ liệu core", "Dùng Fluent API; migration là artifact được commit."],
                  ["Microsoft.AspNetCore.Identity.EntityFrameworkCore", "User, role, password hash", "Authentication and authorization", "Map AppUser and role tables vào SQL Server."],
                  ["Microsoft.AspNetCore.Authentication.JwtBearer", "JWT", "Login, protected endpoint", "Claim role and subject; không tự viết token validation."],
                  ["FluentValidation.AspNetCore", "Request validation", "Create, update, purchase", "Validation syntax tách khỏi controller."],
                  ["MediatR", "Command and query dispatch", "Application layer", "Hữu ích để tổ chức use case; không bắt buộc nếu nhóm muốn service đơn giản."],
                  ["MassTransit and RabbitMQ transport", "Messaging and Saga", "Purchase package, AME result", "Dùng state machine persistence bằng EF Core."],
                  ["Quartz.Extensions.Hosting", "Scheduled jobs", "Expiration, outbox retry, AME retry", "Chỉ bật sau khi database core hoàn chỉnh."],
                  ["Serilog.AspNetCore and sinks", "Structured logging", "All API and worker flows", "Log JSON with correlation and workflow identifiers."],
                  ["AspNetCore.HealthChecks.SqlServer", "Readiness and liveness", "Deployment demo", "Thêm RabbitMQ check nếu package tương thích."],
                  ["Swashbuckle.AspNetCore", "Swagger UI", "API demo", "Document security scheme and sample idempotency header."],
                  ["xUnit, FluentAssertions, Testcontainers optional", "Testing", "Delta, wallet and Saga", "Ưu tiên unit test; dùng integration test cho transaction and message flow."],
                  ["FastAPI, Pydantic, Ollama", "AME phase", "Auto moderation", "Chạy service Python tách biệt; .NET chỉ gọi contract và nhận event result."],
              ], widths=[1.85, 1.45, 1.75, 1.95], font_size=8.9)
    add_text(doc, "Dapper là tùy chọn. Chỉ thêm vào query báo cáo hoặc history projection khi EF Core không cho query dễ đọc hoặc đủ nhanh; không trộn EF và Dapper trong cùng một transaction write nếu nhóm chưa hiểu rõ connection and transaction sharing.")


def add_features(doc):
    heading(doc, "4 Feature backlog và tiêu chí hoàn thành", 1)
    add_table(doc,
              ["Mã", "Feature", "Nghiệp vụ chính", "Công nghệ cần chứng minh", "Ưu tiên"],
              [
                  ["F01", "Authentication and profile", "Đăng ký, đăng nhập, role Seller Moderator Admin.", "Identity, JWT, policy, password hashing.", "Must"],
                  ["F02", "Listing lifecycle", "Tạo draft, sửa, submit, ẩn, public listing.", "Controller, binding, FluentValidation, EF Core, resource policy.", "Must"],
                  ["F03", "Manual moderation", "Moderator approve, reject, reason, audit.", "Role policy, state transition, history event.", "Must"],
                  ["F04", "History Delta", "So sánh trước sau, lưu delta, tái tạo lịch sử.", "JSON serialization, algorithm, unit test, query projection.", "Must"],
                  ["F05", "Wallet and package", "Nạp demo balance, mua Standard VIP Boost, refund.", "Transaction, RowVersion, decimal precision, idempotency.", "Must"],
                  ["F06", "Purchase Saga", "Charge, upgrade, history, outbox, compensation.", "MassTransit, RabbitMQ, persisted state machine, outbox and inbox.", "Should"],
                  ["F07", "Operations", "Expire listing, retry outbox, health and logs.", "Quartz, BackgroundService, Options, Serilog, health checks.", "Should"],
                  ["F08", "Workflow status", "FE theo dõi correlation ID và trạng thái mua gói.", "202 Accepted, idempotent replay, read model.", "Should"],
                  ["F09", "AME text moderation", "Rule check, price check, AI advice, decision result.", "FastAPI, Ollama, async integration, retry and manual fallback.", "Later"],
                  ["F10", "AME media moderation", "Ảnh watermark, sensitive content and video scan.", "Queue worker, storage, computer vision model integration.", "Later"],
              ], widths=[0.45, 1.20, 2.20, 2.25, 0.60], font_size=8.8)
    heading(doc, "4 1 Definition of Done cho một feature", 2)
    add_bullets(doc, [
        "Có request and response DTO, validator, authorization rule và ProblemDetails khi lỗi.",
        "Business rule nằm trong Domain or Application, không copy vào nhiều controller.",
        "Database migration, index và dữ liệu seed cần thiết đã được thêm.",
        "Có log chứa correlation ID; feature write có audit or history phù hợp.",
        "Có unit test cho rule quan trọng; workflow có ít nhất happy path và failure path.",
        "Swagger mô tả endpoint, JWT yêu cầu và header Idempotency-Key khi áp dụng.",
    ])


def add_database(doc):
    heading(doc, "5 Thiết kế database SQL Server", 1)
    add_text(doc, "SQL Server là system of record. Các thao tác ghi luôn đi qua EF Core transaction. RabbitMQ không được xem là nơi lưu trạng thái lâu dài; workflow và idempotency phải có record trong SQL Server.")
    heading(doc, "5 1 Quy ước dữ liệu", 2)
    add_bullets(doc, [
        "Primary key nghiệp vụ dùng BIGINT IDENTITY; message and workflow dùng UNIQUEIDENTIFIER correlation ID.",
        "Money dùng DECIMAL(18,2), diện tích dùng DECIMAL(12,2), timestamp dùng DATETIMEOFFSET.",
        "Enum lưu NVARCHAR có check constraint hoặc converter; dễ đọc khi demo và query database.",
        "Mọi table thay đổi thường xuyên có CreatedAt and UpdatedAt. Listings and Wallets có RowVersion.",
        "JSON chỉ dùng cho flexible snapshot, delta and message payload; cột query chính vẫn là relational column.",
    ])
    heading(doc, "5 2 Nhóm Identity và wallet", 2)
    add_table(doc,
              ["Bảng", "Cột chính", "Khóa và ràng buộc", "Mục đích"],
              [
                  ["Users", "UserId BIGINT, Email, PasswordHash, SecurityStamp, FullName, PhoneNumber, Role, Status, CreatedAt, UpdatedAt", "PK UserId; UQ Email; ApplicationUser kế thừa IdentityUser<long> do IdentityDbContext map.", "Xác thực, bảo mật mật khẩu và hồ sơ người dùng."],
                  ["Roles and UserRoles", "RoleId, Name, UserId", "Identity-managed tables (AspNetRoles, AspNetUserRoles); FK tới Users.", "Quản lý phân quyền Seller, Moderator, Admin."],
                  ["RefreshTokens", "TokenId, UserId, TokenHash, ExpiresAt, RevokedAt", "PK TokenId; index UserId and ExpiresAt.", "Refresh JWT an toàn."],
                  ["Wallets", "WalletId, UserId, MainBalance, PromoBalance, RowVersion, UpdatedAt", "PK WalletId; UQ UserId; FK Users; RowVersion.", "Số dư ví theo người dùng."],
                  ["WalletTransactions", "TransactionId, WalletId, CorrelationId, Type, Amount, BalanceBefore, BalanceAfter, CreatedAt", "PK; FK Wallets; UQ WalletId plus CorrelationId plus Type when CorrelationId exists.", "Ledger bất biến cho nạp, charge, refund."],
                  ["PackageDefinitions", "PackageCode, Name, Price, DurationDays, Kind, IsActive", "PK PackageId; UQ PackageCode; check Price >= 0.", "Bảng giá Standard, VIP and Boost."],
              ], widths=[1.20, 2.30, 2.25, 1.25], font_size=8.7)
    add_text(doc, "Quyết định thiết kế Identity: Nhóm thống nhất tích hợp ASP.NET Core Identity theo trường phái Pragmatic Clean Architecture. Thực thể ApplicationUser ở tầng Domain kế thừa IdentityUser<long> (thông qua package nhẹ Microsoft.Extensions.Identity.Stores) để tận dụng toàn diện UserManager và RoleManager (mã hóa mật khẩu chuẩn PBKDF2, lockout, concurrency stamp). Tầng Infrastructure cấu hình PropNestDbContext kế thừa IdentityDbContext<ApplicationUser, IdentityRole<long>, long> và map bảng chính thành 'Users' với khóa chính 'UserId' để tương thích toàn vẹn với các khóa ngoại từ Listings và Wallets.")
    heading(doc, "5 3 Nhóm listing và moderation", 2)
    add_table(doc,
              ["Bảng", "Cột chính", "Khóa và index", "Mục đích"],
              [
                  ["Listings", "ListingId, UserId, Title, Description, PropertyType, ListingType, Price, Area, City, District, Ward, Address, Status, PackageCode, StartDate, EndDate, RowVersion", "PK ListingId; FK Users; IX UserId Status; IX Status EndDate; IX City District PropertyType.", "Aggregate chính của tin đăng."],
                  ["ListingImages", "ImageId, ListingId, Url, SortOrder, Status, CreatedAt", "PK; FK Listings; UQ ListingId plus SortOrder.", "Metadata ảnh, không lưu binary trong SQL."],
                  ["ListingHistories", "HistoryId, ListingId, ActorUserId, ActionType, ActionDate, IsSnapshot, SnapshotDataJson, DeltaChangesJson, Note", "PK; FK Listings and Users; IX ListingId ActionDate.", "Audit event và delta history."],
                  ["ListingModerationReviews", "ReviewId, ListingId, Source, Decision, ReasonsJson, ReviewerUserId, ReviewedAt", "PK; FK Listings; IX ListingId ReviewedAt.", "Manual or AME review evidence."],
                  ["MarketPriceBenchmarks", "BenchmarkId, CityCode, DistrictCode, PropertyType, ListingType, MinPrice, AvgPrice, MaxPrice, EffectiveFrom", "PK; UQ geography plus property plus type plus effective date.", "Reference price for AME and warning rule."],
                  ["ContentRules", "RuleId, RuleType, Pattern, Severity, IsActive, UpdatedAt", "PK; IX RuleType IsActive.", "Badword and formatting rules controlled by data."],
              ], widths=[1.35, 2.50, 2.15, 1.00], font_size=8.5)
    heading(doc, "5 4 Nhóm workflow và messaging", 2)
    add_table(doc,
              ["Bảng", "Cột chính", "Ràng buộc", "Mục đích"],
              [
                  ["IdempotencyRequests", "IdempotencyKey, UserId, RequestPath, RequestHash, CorrelationId, Status, ResponseCode, ResponseBody, ExpiresAt", "PK RequestId; UQ UserId plus IdempotencyKey; index ExpiresAt.", "Chống request mua gói bị lặp."],
                  ["SagaListingStates", "CorrelationId, CurrentState, ListingId, UserId, PackageCode, ChargeAmount, RetryCount, ErrorMessage, UpdatedAt", "PK CorrelationId; IX CurrentState UpdatedAt.", "Durable state của MassTransit Saga."],
                  ["OutboxMessages", "OutboxId, OccurredOn, Type, PayloadJson, CorrelationId, ProcessedOn, RetryCount, Error", "PK OutboxId; filtered IX ProcessedOn is null.", "Phát integration event đáng tin cậy."],
                  ["InboxMessages", "ConsumerName, MessageId, ReceivedAt, ProcessedAt, Status", "UQ ConsumerName plus MessageId.", "Deduplicate message redelivery tại consumer."],
                  ["WorkflowSteps", "StepId, CorrelationId, StepName, Status, Attempt, StartedAt, FinishedAt, Error", "FK SagaListingStates; IX CorrelationId StepName.", "Read model và evidence để FE xem workflow."],
                  ["AmeModerationResults", "ResultId, ListingId, CorrelationId, Decision, ScoreJson, ReasonsJson, ModelInfo, CheckedAt", "PK; FK Listings; IX ListingId CheckedAt; UQ CorrelationId when one run per submission.", "Kết quả AME có thể audit and replay."],
              ], widths=[1.40, 2.45, 2.05, 1.10], font_size=8.4)
    heading(doc, "5 5 Quan hệ và transaction boundary", 2)
    add_code(doc, [
        "Users 1 to 1 Wallets",
        "Users 1 to many Listings, ListingHistories and ModerationReviews",
        "Listings 1 to many ListingImages, ListingHistories, ModerationReviews and AmeModerationResults",
        "Wallets 1 to many WalletTransactions",
        "SagaListingStates 1 to many WorkflowSteps",
        "A write transaction may include Listing or Wallet change, History, Idempotency record and Outbox message",
    ])
    add_text(doc, "Không update số dư ví bằng phép cộng trừ trực tiếp từ request. Tải Wallet trong transaction, lock or optimistic concurrency, kiểm tra số dư, update wallet, ghi WalletTransaction và commit cùng nhau. Khi có concurrency conflict, trả retry-safe error hoặc thực hiện retry có giới hạn.")


def add_workflows(doc):
    heading(doc, "6 Workflow từng feature", 1)
    heading(doc, "6 1 Tạo và submit listing", 2)
    add_table(doc,
              ["Bước", "Xử lý", "Dữ liệu và kỹ thuật"],
              [
                  ["1", "Seller gọi POST listings với JWT.", "Model binding, FluentValidation, policy Seller."],
                  ["2", "Application tạo Listing Status Draft và ListingHistory snapshot Create.", "EF Core transaction; snapshot JSON chỉ tại mốc cần thiết."],
                  ["3", "Seller cập nhật Draft hoặc Published listing.", "RowVersion chống ghi đè; DeltaEngine so sánh before and after."],
                  ["4", "Seller submit listing.", "Listing chuyển PendingPayment hoặc PendingModeration tùy package."],
                  ["5", "Outbox ghi ListingSubmitted event khi cần moderation or AME.", "Không publish RabbitMQ trước database commit."],
              ], widths=[0.55, 3.55, 3.00])
    heading(doc, "6 2 Cập nhật listing và history delta", 2)
    add_code(doc, [
        "Load Listing with RowVersion",
        "-> Authorize ownership",
        "-> Validate request",
        "-> Clone selected fields before update",
        "-> Apply allowed change",
        "-> DeltaEngine produces [{ field, oldValue, newValue }]",
        "-> Save Listing and ListingHistory in one transaction",
        "-> If material fields changed after publish, set PendingModeration and publish ListingResubmitted",
    ])
    add_text(doc, "API GET listings slash id slash history đọc snapshot gần nhất trước mốc cần xem, áp dụng các delta theo ActionDate và trả BeforeListing, AfterListing and DiffDetails. Unit test phải cover unchanged field, null value, decimal price, nested AttributesJson và concurrent update.")
    heading(doc, "6 3 Duyệt thủ công", 2)
    add_table(doc,
              ["Điều kiện", "Approve", "Reject"],
              [
                  ["Quyền", "Moderator or Admin policy.", "Moderator or Admin policy."],
                  ["Thay đổi listing", "Status Published; StartDate and EndDate được set theo package.", "Status Rejected; không public."],
                  ["Audit", "Review Source Manual, Decision Approved và History Approve.", "Review Source Manual, Decision Rejected, required reason và History Reject."],
                  ["Event", "ListingPublished trong Outbox.", "ListingRejected có lý do cho Seller."],
              ], widths=[1.55, 2.80, 2.75], font_size=9.2)
    heading(doc, "6 4 Mua gói và Saga Orchestrator", 2)
    add_text(doc, "Endpoint POST listings slash id slash purchase-package nhận header Idempotency-Key. API tạo hoặc trả lại IdempotencyRequest, khởi tạo SagaListingState và trả 202 Accepted với correlationId. Worker là orchestration owner.")
    add_table(doc,
              ["State", "Command or event", "Hành động", "Lỗi và bồi hoàn"],
              [
                  ["Started", "PurchasePackageRequested", "Validate listing owner, package and workflow input.", "Invalid input chuyển Failed, không trừ tiền."],
                  ["DeductingWallet", "DeductWalletCommand", "Wallet module trừ PromoBalance trước rồi MainBalance, ghi ledger.", "Insufficient funds chuyển Failed."],
                  ["WalletDeducted", "WalletDeductedEvent", "Saga persist charge success, gửi UpgradeListingCommand.", "Nếu message replay, InboxMessages bỏ qua duplicate."],
                  ["UpgradingListing", "UpgradeListingCommand", "Set package, StartDate and EndDate; ghi history UpgradePackage.", "Lỗi database hoặc business rule chuyển Compensating."],
                  ["RecordingHistory", "ListingUpgradedEvent", "Bảo đảm history and outbox search or moderation event.", "Có retry; nếu unrecoverable, compensation theo policy."],
                  ["Completed", "WorkflowCompleted", "Persist final response and workflow step; FE query status.", "Không chạy lại side effect khi same correlation ID."],
                  ["Compensating", "RefundWalletCommand", "Hoàn đúng amount theo correlation ID, ghi Refund ledger.", "Sau refund chuyển Compensated; nếu refund lỗi, alert and retry."],
              ], widths=[1.05, 1.65, 2.70, 1.70], font_size=8.7)
    add_text(doc, "Tên đúng của mô hình này là Saga Orchestration, không phải choreography. Choreography chỉ phù hợp khi mỗi service tự phản ứng event và không có thành phần trung tâm biết toàn bộ workflow; mô hình đó làm bài demo trạng thái and compensation khó quan sát hơn.")
    heading(doc, "6 5 Idempotency, outbox và inbox", 2)
    add_bullets(doc, [
        "API kiểm tra UserId plus Idempotency-Key. Key cũ cùng RequestHash trả lại response or correlationId cũ; khác RequestHash trả 409 Conflict.",
        "Transaction tạo business state, Saga state and OutboxMessage cùng commit. Dispatcher publish outbox sau commit.",
        "Mỗi consumer ghi InboxMessages trước hoặc cùng xử lý message. MessageId đã tồn tại cho cùng consumer thì ack mà không tạo side effect lần nữa.",
        "WalletTransactions dùng CorrelationId và TransactionType để bảo đảm charge or refund không ghi đôi khi consumer retry.",
    ])
    heading(doc, "6 6 Hết hạn tin", 2)
    add_code(doc, [
        "Quartz trigger every 5 minutes",
        "-> create DI scope",
        "-> query Published listings where EndDate <= now in small batches",
        "-> update status to Expired only if current status is Published",
        "-> create History Expire and Outbox ListingExpired in one transaction",
    ])
    heading(doc, "6 7 AME tự động duyệt tin sau phase core", 2)
    add_table(doc,
              ["Bước", "PropNest", "AME FastAPI", "Kết quả"],
              [
                  ["1", "ListingSubmitted event xuất hiện trong Outbox and RabbitMQ.", "Chưa xử lý.", "Listing PendingModeration."],
                  ["2", "AmeModerationConsumer đọc event và gọi endpoint moderate với correlation ID and timeout.", "Rule engine kiểm tra badword, phone, formatting, price and area; Ollama tư vấn category, spelling and spam.", "Response có decision, reasons, warnings and model metadata."],
                  ["3", "Persist AmeModerationResults và publish AmeModerationCompleted.", "Không truy cập trực tiếp database PropNest.", "Evidence giữ được để audit."],
                  ["4", "Decision handler auto publish, reject or route ManualReview theo policy.", "LLM output không được tự ghi Listing status.", "Listing state được thay đổi bởi .NET domain workflow."],
                  ["5", "Timeout retry theo RetryAt and max attempt; quá giới hạn route ManualReview.", "Có thể offline; fallback là manual review, không default auto approve.", "Không bỏ sót listing."],
              ], widths=[0.45, 2.10, 2.55, 2.00], font_size=8.6)
    add_text(doc, "Khuyến nghị policy AME: rule vi phạm nghiêm trọng như từ cấm hoặc số điện thoại lộ rõ có thể Reject theo rule deterministic; model confidence thấp, price benchmark thiếu hoặc Ollama offline phải ManualReview. LLM chỉ là tín hiệu tư vấn trong decision aggregator, không phải nguồn quyền quyết định duy nhất.")


def add_api_and_operations(doc):
    heading(doc, "7 API chủ đạo và vận hành", 1)
    heading(doc, "7 1 Endpoint đề xuất", 2)
    add_table(doc,
              ["Method and path", "Role", "Kết quả chính", "Ghi chú"],
              [
                  ["POST auth register", "Anonymous", "201 user", "Identity validates password and unique email."],
                  ["POST auth login", "Anonymous", "200 JWT and refresh token", "JWT contains sub and role claim."],
                  ["POST listings", "Seller", "201 draft listing", "Validate and write Create history snapshot."],
                  ["PUT listings slash id", "Owner Seller", "200 updated listing", "Require RowVersion or If Match; write Delta history."],
                  ["POST listings slash id slash submit", "Owner Seller", "202 or 200", "Moves to payment or moderation workflow."],
                  ["POST listings slash id slash purchase package", "Owner Seller", "202 correlationId", "Require Idempotency-Key."],
                  ["GET workflows slash correlationId", "Owner or Admin", "200 workflow status", "Poll while async Saga runs."],
                  ["GET listings slash id slash history", "Owner or Moderator Admin", "200 reconstructed history", "Read model must not mutate data."],
                  ["POST moderation slash listings slash id slash approve", "Moderator Admin", "200 published", "Record manual decision and outbox event."],
                  ["POST moderation slash listings slash id slash reject", "Moderator Admin", "200 rejected", "Reason is required."],
                  ["GET health", "Anonymous or protected", "200 or 503", "Checks app, SQL Server and RabbitMQ readiness."],
              ], widths=[2.25, 1.05, 1.55, 2.25], font_size=8.7)
    heading(doc, "7 2 Cấu hình theo môi trường", 2)
    add_table(doc,
              ["Options class", "Nguồn cấu hình", "Ví dụ field", "Feature"],
              [
                  ["JwtOptions", "appsettings and secret store", "Issuer, Audience, SigningKey, ExpiryMinutes", "Authentication"],
                  ["SqlServerOptions", "connection string", "ConnectionString, CommandTimeout", "EF Core and health check"],
                  ["RabbitMqOptions", "environment variable in deploy", "Host, VirtualHost, UserName", "MassTransit and outbox"],
                  ["AmeOptions", "environment variable", "BaseUrl, TimeoutSeconds, MaxAttempts", "AME integration"],
                  ["ListingPolicyOptions", "appsettings", "DefaultDurationDays, MaterialChangeFields", "Listing lifecycle"],
              ], widths=[1.45, 1.65, 2.45, 1.55], font_size=9)
    add_text(doc, "Không commit JWT signing key, database password or RabbitMQ credentials. Development dùng User Secrets or .env injected by Docker Compose; production dùng secret store. Options phải được validate at startup để lỗi cấu hình lộ ra sớm.")
    heading(doc, "7 3 Logging, health và metric demo", 2)
    add_bullets(doc, [
        "Mỗi request log: correlationId, userId if authenticated, path, method, status code and elapsed time.",
        "Mỗi Saga log: correlationId, state transition, command, attempt, error and compensation action.",
        "Health endpoints: liveness chỉ kiểm tra process; readiness kiểm tra SQL Server, RabbitMQ and optional AME when AME feature is enabled.",
        "Metric tối thiểu: số tin theo status, workflow completed and failed, duplicate idempotency replay, outbox backlog and AME manual fallback count.",
    ])


def add_testing_and_plan(doc):
    heading(doc, "8 Kiểm thử và lộ trình nhóm", 1)
    heading(doc, "8 1 Test strategy", 2)
    add_table(doc,
              ["Loại test", "Đối tượng", "Case tối thiểu"],
              [
                  ["Unit test", "Listing state transition", "Không publish từ Draft; reject yêu cầu reason; hidden không tự expire."],
                  ["Unit test", "DeltaEngine", "Price, null, title, address, AttributesJson, unchanged value and snapshot reconstruction."],
                  ["Unit test", "Wallet policy", "Promo first, insufficient funds, exact refund, duplicate correlation ID."],
                  ["Unit test", "AME DecisionAggregator", "Severe rule reject, low confidence manual review, Ollama unavailable manual review."],
                  ["Integration test", "API and SQL Server", "JWT policy, concurrency conflict, migration and transaction rollback."],
                  ["Integration test", "Saga and messaging", "Happy path, upgrade failure then refund, repeated message no duplicate charge."],
                  ["Manual demo", "Pipeline", "Show correlation ID and exception response before and after middleware order change."],
              ], widths=[1.20, 2.15, 3.75])
    heading(doc, "8 2 Phân chia cho nhóm ba người", 2)
    add_table(doc,
              ["Vai trò", "Module chính", "Deliverable"],
              [
                  ["Thành viên 1", "Listing lifecycle, history delta, authorization ownership", "F02 F03 F04, migrations listing, Delta unit tests, history API."],
                  ["Thành viên 2", "Wallet, package, idempotency, Saga", "F05 F06 F08, workflow status endpoint, refund and concurrency tests."],
                  ["Thành viên 3", "Infrastructure, outbox, worker, operations and later AME bridge", "F07, Docker Compose, logs, health checks, AME integration contract and retry."],
                  ["Cả nhóm", "Architecture decision, API contract, code review and demo", "README, Swagger, migration runbook, end to end scenario and presentation evidence."],
              ], widths=[1.30, 2.55, 3.25])
    heading(doc, "8 3 Sprint roadmap", 2)
    add_table(doc,
              ["Sprint", "Mục tiêu", "Demo kết thúc sprint"],
              [
                  ["1", "Solution, Identity, SQL Server, migrations, listing draft CRUD, pipeline and logging.", "Seller tạo draft; Swagger JWT; migration chạy từ database rỗng."],
                  ["2", "Listing lifecycle, manual moderation, History Delta and tests.", "Moderator publish or reject; xem before and after history."],
                  ["3", "Wallet, package, idempotency and transaction concurrency.", "Mua package demo; retry HTTP cùng key không bị trừ hai lần."],
                  ["4", "RabbitMQ, persisted Saga, outbox, compensation and workflow status.", "Giả lập UpgradeListing lỗi; wallet được refund and Saga Compensated."],
                  ["5", "Quartz, health, integration tests, Docker and report evidence.", "Listing expire tự động; health reports readiness; deployment local one command."],
                  ["6 Later", "AME FastAPI, Ollama and moderation decision integration.", "Tin submit được rule and AI check, sau đó auto publish or manual review."],
              ], widths=[0.75, 3.55, 2.80])
    heading(doc, "8 4 Kịch bản demo cuối kỳ", 2)
    add_bullets(doc, [
        "Seller đăng nhập, tạo Draft, cập nhật giá và xem history delta; chứng minh Delta nhỏ hơn snapshot bằng sample payload.",
        "Seller gửi request mua VIP hai lần với cùng Idempotency-Key; response trả cùng correlationId và database chỉ có một Charge ledger.",
        "Giả lập UpgradeListing failure; quan sát Saga state transition, Refund ledger, outbox and structured logs theo correlation ID.",
        "Moderator duyệt manual hoặc AME trả ManualReview; listing chỉ Published sau quyết định của domain workflow.",
        "Đảo middleware có kiểm soát trong môi trường demo để giải thích correlation ID và exception handling pipeline.",
    ])


def add_ame_guidance(doc):
    heading(doc, "9 Hướng phát triển AME", 1)
    add_text(doc, "AME là Auto Moderation Engine được tích hợp sau khi phần lõi hoàn thành. Mục đích là giảm công việc duyệt lặp lại, nhưng không được làm mất khả năng giải thích hoặc audit quyết định. Tài liệu kiến trúc AME tham khảo gợi ý FastAPI, rule engine và Ollama; blueprint này điều chỉnh nó để phù hợp phạm vi học tập và hệ thống PropNest.")
    heading(doc, "9 1 Phiên bản AME nên xây trước", 2)
    add_table(doc,
              ["Năng lực", "Cách hiện thực", "Kết quả moderation"],
              [
                  ["Badword and phone", "Regex và ContentRules từ SQL Server or AME config.", "Reject khi vi phạm nghiêm trọng."],
                  ["Formatting", "Rule all caps, spam hashtag, abnormal whitespace and unsupported character.", "Warning hoặc ManualReview tùy severity."],
                  ["Price and area", "Tính price per m2 và so MarketPriceBenchmarks theo khu vực and property type.", "ManualReview nếu dữ liệu benchmark thiếu; Reject only with deterministic threshold."],
                  ["Category and spelling", "Ollama local model trả JSON schema strict; validate response bằng Pydantic.", "AI advice; low confidence chuyển ManualReview."],
                  ["Decision aggregation", "Kết hợp hard rules, benchmark, model confidence and policy thresholds.", "Approved, ManualReview or Rejected with reason codes."],
              ], widths=[1.40, 3.45, 2.25], font_size=9)
    heading(doc, "9 2 Contract và bảo mật tích hợp", 2)
    add_bullets(doc, [
        "PropNest gửi listingId, correlationId, title, description, selected property type, location, price, area and version. Không gửi JWT signing key, password or wallet data.",
        "AME trả decision code, reason codes, warnings, category confidence, price score, spelling suggestions, rule version and model identifier.",
        "PropNest validate contract response and persist raw normalized result into AmeModerationResults. LLM free text không dùng trực tiếp để render HTML or build SQL.",
        "Timeout phải ngắn và retry bất đồng bộ. Khi AME unavailable, status là ManualReview, không auto publish vì fallback default true.",
        "Image and video checks chỉ thêm khi có storage strategy and queue worker rõ ràng; không nhét binary media vào database listing.",
    ])
    heading(doc, "9 3 Điều kiện để bắt đầu AME", 2)
    add_bullets(doc, [
        "F01 đến F05 chạy ổn định; một listing có trạng thái moderation, history and owner authorization đúng.",
        "Outbox và retry đã chứng minh không mất ListingSubmitted event.",
        "Nhóm có sample data benchmark, danh sách rule, tiêu chí ManualReview và bộ test accepted or rejected rõ ràng.",
        "AME chạy local bằng Docker Compose hoặc runbook riêng; Ollama model không phải dependency bắt buộc để khởi động API core.",
    ])


def add_final_checklist(doc):
    heading(doc, "10 Checklist chốt trước khi code", 1)
    add_bullets(doc, [
        "Chốt Target Framework duy nhất cho toàn solution. Không trộn .NET 8 và .NET 9 trong cùng project graph.",
        "Chốt enum listing, moderation, workflow and transaction type trước migration đầu tiên.",
        "Chốt API response envelope and ProblemDetails; Swagger là contract chung cho FE and BE.",
        "Chốt Idempotency-Key policy: scope by UserId, TTL, hash payload and behavior when same key has different payload.",
        "Chốt một Saga happy path và một compensation path trước khi thêm feature AME.",
        "Chỉ bật RabbitMQ and Quartz khi đã có health check, log and repeatable Docker Compose runbook.",
        "Mỗi milestone phải có demo script, migration command, sample account and test evidence.",
    ])
    add_text(doc, "Kết luận: hướng đi phù hợp nhất là xây phần lõi PropNest như Modular Monolith có workflow worker riêng, rồi dùng AME như một dịch vụ bổ sung theo event. Kiến trúc này cho phép nhóm chứng minh pipeline, DI, EF Core, authorization, concurrency, idempotency, Saga, outbox và background jobs bằng các nghiệp vụ có thể demo end to end.", space_after=0)


def main():
    doc = Document()
    configure_document(doc)
    title_page(doc)
    add_overview(doc)
    add_architecture(doc)
    add_dotnet_practices(doc)
    add_features(doc)
    add_database(doc)
    add_workflows(doc)
    add_api_and_operations(doc)
    add_testing_and_plan(doc)
    add_ame_guidance(doc)
    add_final_checklist(doc)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
